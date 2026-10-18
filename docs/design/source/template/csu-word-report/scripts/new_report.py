#!/usr/bin/env python3
"""Create a new, editable CSU-style report without replacing an existing file.

The document uses Word section numbering, heading numbering, TOC, SEQ and REF
fields. Field evaluation belongs to Word or another compatible word processor;
this module does not fabricate a table of contents or predicted page numbers.
"""
from __future__ import annotations

import argparse
import json
import math
import re
import sys
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Iterable

from docx import Document
from docx.document import Document as DocumentObject
from docx.enum.section import WD_SECTION_START
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image, ImageFont


SKILL_DIR = Path(__file__).resolve().parents[1]
ASSET_DIR = SKILL_DIR / "assets"
TEXT_WIDTH_CM = 15.6
BODY_FONT = "SimSun"
STATIC_NOTO_REGULAR = Path("C:/Windows/Fonts/Noto Sans SC (TrueType).otf")
STATIC_NOTO_BOLD = Path("C:/Windows/Fonts/Noto Sans SC Bold (TrueType).otf")
HEADING_FONT = "Noto Sans SC" if STATIC_NOTO_REGULAR.is_file() and STATIC_NOTO_BOLD.is_file() else "SimHei"
BOOKMARK_PATTERN = re.compile(r"^[A-Za-z][A-Za-z0-9_]{0,39}$")
REF_PATTERN = re.compile(r"\[\[REF:([A-Za-z][A-Za-z0-9_]{0,39})\]\]")
SHORT_ID_PATTERN = re.compile(r"^[A-Za-z][A-Za-z0-9]*(?:[-_][A-Za-z0-9]+)*$")


@dataclass(frozen=True)
class ReportMetadata:
    title: str
    report_type: str = ""
    subtitle: str = ""
    college: str = ""
    class_name: str = ""
    student_id: str = ""
    author: str = ""
    advisor: str = ""
    header_title: str = ""


def xml_child(parent: Any, tag: str, **attributes: Any) -> Any:
    """Get or add a namespaced OOXML child and set its attributes."""
    element = parent.find(qn(tag))
    if element is None:
        element = OxmlElement(tag)
        parent.append(element)
    for key, value in attributes.items():
        element.set(qn(key if ":" in key else f"w:{key}"), str(value))
    return element


def set_font(target: Any, *, name: str = BODY_FONT, size: float = 12,
             bold: bool = False) -> None:
    """Apply all Word font slots, preventing an East Asian font fallback."""
    target.font.name = name
    target.font.size = Pt(size)
    target.font.bold = bold
    target.font.color.rgb = RGBColor(0, 0, 0)
    target.font.underline = False
    rpr = target._element.get_or_add_rPr()
    fonts = xml_child(rpr, "w:rFonts")
    for slot in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{slot}"), name)
    for slot in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "csTheme"):
        fonts.attrib.pop(qn(f"w:{slot}"), None)
    xml_child(rpr, "w:lang", val="zh-CN", eastAsia="zh-CN")
    xml_child(rpr, "w:szCs", val=round(size * 2))
    xml_child(rpr, "w:bCs", val="1" if bold else "0")
    color = xml_child(rpr, "w:color", val="000000")
    for attribute in ("themeColor", "themeShade", "themeTint"):
        color.attrib.pop(qn(f"w:{attribute}"), None)


def format_paragraph(paragraph: Any, *, first_line: float = 24,
                     line: float = 26, before: float = 0, after: float = 0,
                     alignment: Any = None, keep_next: bool = False,
                     keep_together: bool = False, widow_control: bool = True) -> None:
    fmt = paragraph.paragraph_format
    fmt.first_line_indent = Pt(first_line)
    fmt.left_indent = Pt(0)
    fmt.right_indent = Pt(0)
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = Pt(line)
    fmt.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    fmt.keep_with_next = keep_next
    fmt.keep_together = keep_together
    fmt.widow_control = widow_control
    fmt.page_break_before = False
    if alignment is not None:
        paragraph.alignment = alignment
    xml_child(paragraph._element.get_or_add_pPr(), "w:snapToGrid", val="0")


def add_field(paragraph: Any, instruction: str, result: str = "", *,
              size: float = 12, bold: bool = False,
              font: str = BODY_FONT) -> None:
    """Append a real complex field with an optional, honest cached result."""
    begin = paragraph.add_run()
    set_font(begin, name=font, size=size, bold=bold)
    xml_child(begin._element, "w:fldChar", fldCharType="begin", dirty="true")
    code = paragraph.add_run()
    set_font(code, name=font, size=size, bold=bold)
    instruction_element = OxmlElement("w:instrText")
    instruction_element.set(qn("xml:space"), "preserve")
    instruction_element.text = f" {instruction} "
    code._element.append(instruction_element)
    separate = paragraph.add_run()
    set_font(separate, name=font, size=size, bold=bold)
    xml_child(separate._element, "w:fldChar", fldCharType="separate")
    cached = paragraph.add_run(result)
    set_font(cached, name=font, size=size, bold=bold)
    end = paragraph.add_run()
    set_font(end, name=font, size=size, bold=bold)
    xml_child(end._element, "w:fldChar", fldCharType="end")


def add_text_with_refs(paragraph: Any, text: str, bookmarks: dict[str, int],
                       *, size: float = 12, bold: bool = False,
                       font: str = BODY_FONT) -> None:
    """Turn [[REF:bookmark]] markers into linked REF fields."""
    position = 0
    for match in REF_PATTERN.finditer(text):
        if match.start() > position:
            set_font(paragraph.add_run(text[position:match.start()]), name=font, size=size, bold=bold)
        name = match.group(1)
        if name not in bookmarks:
            raise ValueError(f"Reference uses an undefined figure bookmark: {name}")
        add_field(paragraph, f"REF {name} \\h", str(bookmarks[name]), size=size, bold=bold, font=font)
        position = match.end()
    if position < len(text) or not paragraph.runs:
        set_font(paragraph.add_run(text[position:]), name=font, size=size, bold=bold)


def configure_styles(document: DocumentObject) -> None:
    styles = document.styles
    normal = styles["Normal"]
    set_font(normal)
    normal.paragraph_format.first_line_indent = Pt(24)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = Pt(26)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    normal.paragraph_format.widow_control = True
    normal.paragraph_format.keep_with_next = False
    xml_child(normal._element.get_or_add_pPr(), "w:snapToGrid", val="0")
    cover_title = styles["Title"]
    cover_title.base_style = normal
    set_font(cover_title, name=HEADING_FONT, size=19, bold=True)
    cover_title.paragraph_format.first_line_indent = Pt(0)
    cover_title.paragraph_format.space_before = Pt(0)
    cover_title.paragraph_format.space_after = Pt(0)
    cover_title.paragraph_format.line_spacing = Pt(40.3)
    cover_title.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    cover_title.paragraph_format.keep_together = True
    title_ppr = cover_title._element.get_or_add_pPr()
    for border in list(title_ppr.findall(qn("w:pBdr"))):
        title_ppr.remove(border)
    xml_child(title_ppr, "w:snapToGrid", val="0")
    for level, size, before, after in ((1, 15, 18, 8), (2, 14, 14, 6), (3, 12, 10, 4)):
        style = styles[f"Heading {level}"]
        set_font(style, name=HEADING_FONT, size=size, bold=True)
        fmt = style.paragraph_format
        fmt.first_line_indent = Pt(0)
        fmt.left_indent = Pt(0)
        fmt.space_before = Pt(before)
        fmt.space_after = Pt(after)
        fmt.line_spacing = Pt(26)
        fmt.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        fmt.keep_with_next = True
        fmt.keep_together = True
        fmt.widow_control = True
        fmt.page_break_before = False
        ppr = style._element.get_or_add_pPr()
        xml_child(ppr, "w:outlineLvl", val=level - 1)
        xml_child(ppr, "w:snapToGrid", val="0")
    custom_styles = (
        ("Report Title", 18, HEADING_FONT, True, 0, 10),
        ("Front Matter Title", 15, HEADING_FONT, True, 0, 18),
        ("Report Reference", 12, BODY_FONT, False, 0, 0),
    )
    for name, size, font, bold, before, after in custom_styles:
        style = styles[name] if name in styles else styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        style.base_style = normal
        set_font(style, name=font, size=size, bold=bold)
        fmt = style.paragraph_format
        fmt.first_line_indent = Pt(0)
        fmt.space_before = Pt(before)
        fmt.space_after = Pt(after)
        fmt.line_spacing = Pt(26 if name != "Report Title" else 36.4)
        fmt.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        fmt.keep_with_next = name != "Report Reference"
        fmt.keep_together = True
        fmt.widow_control = True
        xml_child(style._element.get_or_add_pPr(), "w:outlineLvl", val="9")
        xml_child(style._element.get_or_add_pPr(), "w:snapToGrid", val="0")
    caption = styles["Caption"]
    caption.base_style = normal
    set_font(caption, size=11)
    caption.font.italic = False
    caption.paragraph_format.first_line_indent = Pt(0)
    caption.paragraph_format.line_spacing = Pt(18)
    caption.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    caption.paragraph_format.keep_together = True
    caption.paragraph_format.keep_with_next = False
    xml_child(caption._element.get_or_add_pPr(), "w:snapToGrid", val="0")
    for name, indent in (("TOC 1", 0), ("TOC 2", 12), ("Table of Figures", 0)):
        style = styles[name] if name in styles else styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        style.base_style = normal
        set_font(style)
        fmt = style.paragraph_format
        fmt.first_line_indent = Pt(0)
        fmt.left_indent = Pt(indent)
        fmt.space_before = Pt(0)
        fmt.space_after = Pt(0)
        fmt.line_spacing = Pt(26)
        fmt.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        fmt.keep_with_next = False
        fmt.keep_together = True
        xml_child(style._element.get_or_add_pPr(), "w:snapToGrid", val="0")
    for name in ("Hyperlink", "FollowedHyperlink"):
        # Use Word's built-in styles so updating a TOC does not replace a custom
        # look-alike with the default blue Hyperlink style.
        style = styles[name] if name in styles else styles.add_style(
            name, WD_STYLE_TYPE.CHARACTER, builtin=True)
        set_font(style)
        color = xml_child(style._element.get_or_add_rPr(), "w:color", val="000000")
        for attribute in ("themeColor", "themeTint", "themeShade"):
            color.attrib.pop(qn("w:" + attribute), None)
    for name in ("Header", "Footer"):
        set_font(styles[name], size=10)
        styles[name].paragraph_format.first_line_indent = Pt(0)
        styles[name].paragraph_format.space_after = Pt(0)
        xml_child(styles[name]._element.get_or_add_pPr(), "w:snapToGrid", val="0")


def add_numbering(document: DocumentObject, *, appendix: bool = False) -> int:
    """Add a genuine, editable three-level Word numbering definition."""
    root = document.part.numbering_part.element
    abstract_ids = [int(el.get(qn("w:abstractNumId"))) for el in root.findall(qn("w:abstractNum"))]
    num_ids = [int(el.get(qn("w:numId"))) for el in root.findall(qn("w:num"))]
    abstract_id = max(abstract_ids, default=-1) + 1
    num_id = max(num_ids, default=0) + 1
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    xml_child(abstract, "w:multiLevelType", val="multilevel")
    for level in range(3):
        lvl = OxmlElement("w:lvl")
        lvl.set(qn("w:ilvl"), str(level))
        xml_child(lvl, "w:start", val="1")
        xml_child(lvl, "w:numFmt", val="upperLetter" if appendix and level == 0 else "decimal")
        # Appendix paragraphs use the same outline styles as ordinary chapters.
        # Keep the ordinary outline/style linkage unique to its numbering scheme.
        if not appendix:
            xml_child(lvl, "w:pStyle", val=document.styles[f"Heading {level + 1}"].style_id)
        text = ".".join(f"%{n}" for n in range(1, level + 2))
        if level == 0:
            text = f"{text}."
        xml_child(lvl, "w:lvlText", val=text)
        xml_child(lvl, "w:suff", val="space")
        xml_child(lvl, "w:lvlJc", val="left")
        ppr = xml_child(lvl, "w:pPr")
        xml_child(ppr, "w:ind", left="0", hanging="0")
        rpr = xml_child(lvl, "w:rPr")
        fonts = xml_child(rpr, "w:rFonts")
        for slot in ("ascii", "hAnsi", "eastAsia", "cs"):
            fonts.set(qn(f"w:{slot}"), HEADING_FONT)
        xml_child(rpr, "w:b", val="1")
        xml_child(rpr, "w:color", val="000000")
        xml_child(rpr, "w:sz", val=(30, 28, 24)[level])
        abstract.append(lvl)
    # abstractNum elements precede num elements in the OOXML content model.
    first_num = root.find(qn("w:num"))
    if first_num is None:
        root.append(abstract)
    else:
        root.insert(root.index(first_num), abstract)
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    xml_child(num, "w:abstractNumId", val=abstract_id)
    root.append(num)
    return num_id


def apply_numbering(paragraph: Any, num_id: int, level: int) -> None:
    numpr = xml_child(paragraph._element.get_or_add_pPr(), "w:numPr")
    xml_child(numpr, "w:ilvl", val=level - 1)
    xml_child(numpr, "w:numId", val=num_id)


def configure_section(section: Any, *, page_format: str | None = None) -> None:
    section.start_type = WD_SECTION_START.NEW_PAGE
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(2.7)
    section.right_margin = Cm(2.7)
    section.header_distance = Cm(1.0)
    section.footer_distance = Cm(1.2)
    section.different_first_page_header_footer = False
    # Default/no grid plus explicit snapToGrid=0 on paragraph styles.
    for grid in list(section._sectPr.findall(qn("w:docGrid"))):
        section._sectPr.remove(grid)
    if page_format:
        xml_child(section._sectPr, "w:pgNumType", fmt=page_format, start="1")
    else:
        existing = section._sectPr.find(qn("w:pgNumType"))
        if existing is not None:
            section._sectPr.remove(existing)


def _clear_container(container: Any) -> None:
    for child in list(container._element):
        container._element.remove(child)


def _cell_margins(cell: Any, top: int = 0, bottom: int = 0,
                  start: int = 0, end: int = 0) -> None:
    margins = xml_child(cell._tc.get_or_add_tcPr(), "w:tcMar")
    for name, width in (("top", top), ("bottom", bottom), ("start", start), ("end", end)):
        xml_child(margins, f"w:{name}", w=width, type="dxa")


def _border(parent: Any, edge: str, *, size: int = 0, visible: bool = False) -> None:
    xml_child(parent, f"w:{edge}", val="single" if visible else "nil",
              sz=size, color="000000", space="0")


def _set_table_widths(table: Any, widths_cm: Iterable[float]) -> None:
    widths = list(widths_cm)
    table.autofit = False
    for column, width in zip(table.columns, widths):
        column.width = Cm(width)
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = Cm(width)
    xml_child(table._tbl.tblPr, "w:tblW", w=round(sum(widths) * 567), type="dxa")
    xml_child(table._tbl.tblPr, "w:tblLayout", type="fixed")


def text_width_points(text: str, size: float, *, bold: bool = False) -> float:
    """Measure a header line using the available static font, or a safe estimate."""
    path = STATIC_NOTO_BOLD if bold else STATIC_NOTO_REGULAR
    if path.is_file():
        font = ImageFont.truetype(str(path), 400)
        return font.getlength(text) * size / 400
    return sum(1.0 if ord(ch) > 255 else 0.6 for ch in text) * size


def _fit_font(text: str, width_cm: float, nominal: float) -> float:
    measured = text_width_points(text, nominal)
    available = width_cm * 72 / 2.54
    return round(min(nominal, nominal * available / measured if measured else nominal) * 2) / 2


def _header_line(cell: Any, text: str, *, width_cm: float, nominal: float,
                 align: Any, first: bool) -> None:
    text = " ".join(text.split())
    paragraph = cell.paragraphs[0] if first else cell.add_paragraph()
    format_paragraph(paragraph, first_line=0, line=13, alignment=align)
    size = max(1, _fit_font(text, width_cm, nominal))
    run = paragraph.add_run(text)
    set_font(run, name=HEADING_FONT, size=size)
    # Prevent a long college/name line from adding a third header line.
    xml_child(cell._tc.get_or_add_tcPr(), "w:noWrap")


def create_running_header_footer(section: Any, metadata: ReportMetadata) -> None:
    section.header.is_linked_to_previous = False
    section.footer.is_linked_to_previous = False
    _clear_container(section.header)
    _clear_container(section.footer)
    table = section.header.add_table(rows=1, cols=3, width=Cm(TEXT_WIDTH_CM))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    widths = (3.9, 7.8, 3.9)
    _set_table_widths(table, widths)
    borders = xml_child(table._tbl.tblPr, "w:tblBorders")
    for edge in ("top", "left", "right", "insideH", "insideV"):
        _border(borders, edge)
    # Word stores border thickness in eighths of a point: 5/8 is nearest 0.6pt.
    _border(borders, "bottom", size=5, visible=True)
    row = table.rows[0]
    row.height = Cm(1.06)
    row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
    for cell in row.cells:
        _cell_margins(cell, bottom=25)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    logo_paragraph = row.cells[0].paragraphs[0]
    format_paragraph(logo_paragraph, first_line=0, line=1, alignment=WD_ALIGN_PARAGRAPH.LEFT)
    logo_paragraph.paragraph_format.line_spacing = 1
    logo_paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    logo_paragraph.add_run().add_picture(str(ASSET_DIR / "csu_seal.png"), height=Cm(0.95))
    _header_line(row.cells[1], metadata.report_type or "中南大学报告", width_cm=7.5,
                 nominal=10, align=WD_ALIGN_PARAGRAPH.CENTER, first=True)
    header_title = metadata.header_title or " ".join(metadata.title.split())[:18]
    _header_line(row.cells[1], header_title, width_cm=7.5,
                 nominal=9, align=WD_ALIGN_PARAGRAPH.CENTER, first=False)
    _header_line(row.cells[2], metadata.college, width_cm=3.65, nominal=9.5,
                 align=WD_ALIGN_PARAGRAPH.RIGHT, first=True)
    identity = "　".join(value for value in (metadata.author, metadata.student_id) if value)
    _header_line(row.cells[2], identity, width_cm=3.65, nominal=9.5,
                 align=WD_ALIGN_PARAGRAPH.RIGHT, first=False)
    footer = section.footer.add_paragraph(style="Footer")
    format_paragraph(footer, first_line=0, line=20, alignment=WD_ALIGN_PARAGRAPH.CENTER)
    set_font(footer.add_run("第"), size=10)
    add_field(footer, "PAGE", "I", size=10)
    set_font(footer.add_run("页"), size=10)


def _blank_space(document: DocumentObject, points: float) -> None:
    if points <= 0:
        return
    paragraph = document.add_paragraph()
    format_paragraph(paragraph, first_line=0, line=max(points, 1), widow_control=False)
    set_font(paragraph.add_run(""), size=1)


def _wrapped_lines(text: str, size: float, width_cm: float) -> int:
    capacity = width_cm * 72 / 2.54
    return sum(max(1, math.ceil(text_width_points(line, size, bold=True) / capacity)) for line in text.split("\n"))


def create_cover(document: DocumentObject, metadata: ReportMetadata) -> None:
    """Add the editable one-page cover, including a borderless signature table."""
    cover = document.sections[0]
    configure_section(cover)
    cover.header.is_linked_to_previous = False
    cover.footer.is_linked_to_previous = False
    # Explicitly blank definitions ensure no cover header/footer appears later.
    for container in (cover.header, cover.footer):
        _clear_container(container)
        paragraph = container.add_paragraph()
        format_paragraph(paragraph, first_line=0, line=1)
    with Image.open(ASSET_DIR / "csu_logo_full.png") as image:
        image_height_cm = 14.35 * image.height / image.width
    used = (0.3 + image_height_cm + 1.2) * 72 / 2.54
    logo = document.add_paragraph()
    format_paragraph(logo, first_line=0, before=Cm(0.3).pt, after=Cm(1.2).pt,
                     alignment=WD_ALIGN_PARAGRAPH.CENTER)
    logo.paragraph_format.line_spacing = 1
    logo.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    logo.add_run().add_picture(str(ASSET_DIR / "csu_logo_full.png"), width=Cm(14.35))
    if metadata.report_type:
        paragraph = document.add_paragraph()
        format_paragraph(paragraph, first_line=0, line=39, after=Cm(1.6).pt,
                         alignment=WD_ALIGN_PARAGRAPH.CENTER, keep_next=True)
        set_font(paragraph.add_run(metadata.report_type), name=HEADING_FONT, size=20, bold=True)
        used += 39 * _wrapped_lines(metadata.report_type, 20, TEXT_WIDTH_CM) + Cm(1.6).pt
    paragraph = document.add_paragraph(style="Title")
    format_paragraph(paragraph, first_line=0, line=40.3, alignment=WD_ALIGN_PARAGRAPH.CENTER,
                     keep_next=bool(metadata.subtitle), keep_together=True)
    set_font(paragraph.add_run(metadata.title), name=HEADING_FONT, size=19, bold=True)
    used += 40.3 * _wrapped_lines(metadata.title, 19, TEXT_WIDTH_CM)
    if metadata.subtitle:
        paragraph = document.add_paragraph()
        format_paragraph(paragraph, first_line=0, line=31.2, before=Cm(0.9).pt,
                         alignment=WD_ALIGN_PARAGRAPH.CENTER, keep_together=True)
        set_font(paragraph.add_run(metadata.subtitle), size=14)
        used += Cm(0.9).pt + 31.2 * _wrapped_lines(metadata.subtitle, 14, TEXT_WIDTH_CM)
    signature_height = Cm(8.0).pt
    signature_top = max(Cm(12.7).pt, used + Cm(0.7).pt)
    university_top = Cm(23.0).pt
    if signature_top + signature_height + Cm(0.4).pt > university_top:
        raise ValueError("Cover title/type/subtitle is too long for a single A4 cover; shorten those fields.")
    _blank_space(document, signature_top - used)
    table = document.add_table(rows=5, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    _set_table_widths(table, (3.1, 0.55, 6.4))
    borders = xml_child(table._tbl.tblPr, "w:tblBorders")
    for edge in ("top", "bottom", "left", "right", "insideH", "insideV"):
        _border(borders, edge)
    values = (("学院：", metadata.college), ("班级：", metadata.class_name),
              ("学号：", metadata.student_id), ("姓名：", metadata.author),
              ("指导教师：", metadata.advisor))
    for row, (label, value) in zip(table.rows, values):
        row.height = Cm(1.6)
        row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
        xml_child(row._tr.get_or_add_trPr(), "w:cantSplit")
        for cell in row.cells:
            _cell_margins(cell, top=0, bottom=0, start=0, end=0)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        label_paragraph = row.cells[0].paragraphs[0]
        format_paragraph(label_paragraph, first_line=0, line=26, alignment=WD_ALIGN_PARAGRAPH.RIGHT)
        set_font(label_paragraph.add_run(label), name=HEADING_FONT, size=12, bold=True)
        spacer = row.cells[1].paragraphs[0]
        format_paragraph(spacer, first_line=0, line=26)
        value_paragraph = row.cells[2].paragraphs[0]
        format_paragraph(value_paragraph, first_line=0, line=26, alignment=WD_ALIGN_PARAGRAPH.CENTER)
        value_size = max(7, _fit_font(value, 6.3, 12))
        set_font(value_paragraph.add_run(value), size=value_size)
        xml_child(row.cells[2]._tc.get_or_add_tcPr(), "w:noWrap")
        value_borders = xml_child(value_paragraph._element.get_or_add_pPr(), "w:pBdr")
        _border(value_borders, "bottom", size=4, visible=True)
    # A normal paragraph after the table provides adjustable vertical whitespace.
    _blank_space(document, university_top - signature_top - signature_height)
    paragraph = document.add_paragraph()
    format_paragraph(paragraph, first_line=0, line=26, alignment=WD_ALIGN_PARAGRAPH.CENTER)
    set_font(paragraph.add_run("CENTRAL SOUTH UNIVERSITY"), name=HEADING_FONT, size=12)


def create_front_matter(document: DocumentObject, metadata: ReportMetadata,
                        *, include_figure_list: bool) -> None:
    section = document.add_section(WD_SECTION_START.NEW_PAGE)
    configure_section(section, page_format="upperRoman")
    create_running_header_footer(section, metadata)
    title = document.add_paragraph("目　录", style="Front Matter Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    field_paragraph = document.add_paragraph()
    format_paragraph(field_paragraph, first_line=0)
    add_field(field_paragraph, 'TOC \\o "1-2" \\h \\z \\u')
    if include_figure_list:
        page_break = document.add_paragraph()
        format_paragraph(page_break, first_line=0, line=1)
        page_break.add_run().add_break(WD_BREAK.PAGE)
        title = document.add_paragraph("图　录", style="Front Matter Title")
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        field_paragraph = document.add_paragraph()
        format_paragraph(field_paragraph, first_line=0)
        add_field(field_paragraph, 'TOC \\h \\z \\c "Figure"')


def add_caption(document: DocumentObject, caption: str, *, label: str,
                sequence: str, number: int, bookmark: str | None = None,
                bookmark_id: int = 0, keep_next: bool = False) -> Any:
    paragraph = document.add_paragraph(style="Caption")
    format_paragraph(paragraph, first_line=0, line=18, before=3, after=6,
                     alignment=WD_ALIGN_PARAGRAPH.CENTER,
                     keep_next=keep_next, keep_together=True)
    set_font(paragraph.add_run(f"{label} "), size=11, bold=True)
    if bookmark:
        start = OxmlElement("w:bookmarkStart")
        start.set(qn("w:id"), str(bookmark_id))
        start.set(qn("w:name"), bookmark)
        paragraph._element.append(start)
    add_field(paragraph, f"SEQ {sequence} \\* ARABIC", str(number), size=11, bold=True)
    if bookmark:
        end = OxmlElement("w:bookmarkEnd")
        end.set(qn("w:id"), str(bookmark_id))
        paragraph._element.append(end)
    set_font(paragraph.add_run(f"　{caption}"), size=11)
    return paragraph


def add_three_line_table(document: DocumentObject, headers: list[str], rows: list[list[str]],
                         widths_cm: list[float] | None = None, *,
                         caption: str = "", number: int = 1) -> Any:
    """Add an editable table; repeat the header and let long tables paginate."""
    if caption:
        add_caption(document, caption, label="表", sequence="Table", number=number, keep_next=True)
    widths = widths_cm or [TEXT_WIDTH_CM / len(headers)] * len(headers)
    table = document.add_table(rows=1 + len(rows), cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    _set_table_widths(table, widths)
    borders = xml_child(table._tbl.tblPr, "w:tblBorders")
    for edge in ("left", "right", "insideH", "insideV"):
        _border(borders, edge)
    _border(borders, "top", size=7, visible=True)
    _border(borders, "bottom", size=7, visible=True)
    xml_child(table.rows[0]._tr.get_or_add_trPr(), "w:tblHeader")
    for row_index, (row, data) in enumerate(zip(table.rows, [headers, *rows])):
        xml_child(row._tr.get_or_add_trPr(), "w:cantSplit")
        for cell, text in zip(row.cells, data):
            _cell_margins(cell, top=45, bottom=45, start=85, end=85)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            paragraph = cell.paragraphs[0]
            format_paragraph(paragraph, first_line=0, line=26,
                             keep_next=(row_index == 0), keep_together=True)
            set_font(paragraph.add_run(str(text)), name=HEADING_FONT if row_index == 0 else BODY_FONT,
                     size=12, bold=(row_index == 0))
            if SHORT_ID_PATTERN.fullmatch(str(text)) and len(str(text)) <= 30:
                xml_child(cell._tc.get_or_add_tcPr(), "w:noWrap")
            cell_borders = xml_child(cell._tc.get_or_add_tcPr(), "w:tcBorders")
            for edge in ("left", "right"):
                _border(cell_borders, edge)
            if row_index == 0:
                _border(cell_borders, "top", size=7, visible=True)
                _border(cell_borders, "bottom", size=4, visible=True)
            elif row_index == len(table.rows) - 1:
                _border(cell_borders, "bottom", size=7, visible=True)
            else:
                _border(cell_borders, "bottom")
    return table


def add_figure(document: DocumentObject, path: Path, caption: str, *,
               number: int, bookmark: str, bookmark_id: int,
               width_cm: float | None = None) -> None:
    with Image.open(path) as image:
        ratio = image.height / image.width
    width = min(width_cm or 14.0, TEXT_WIDTH_CM)
    # Preserve aspect ratio and leave room for a caption on a portrait page.
    width = min(width, 21.5 / ratio)
    paragraph = document.add_paragraph()
    format_paragraph(paragraph, first_line=0, before=6,
                     alignment=WD_ALIGN_PARAGRAPH.CENTER,
                     keep_next=True, keep_together=True)
    # Fixed 26pt text spacing would crop an inline image.
    paragraph.paragraph_format.line_spacing = 1
    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    paragraph.add_run().add_picture(str(path), width=Cm(width))
    add_caption(document, caption, label="图", sequence="Figure", number=number,
                bookmark=bookmark, bookmark_id=bookmark_id)


def default_blocks() -> list[dict[str, Any]]:
    return [
        {"type": "paragraph", "text": "此底稿展示封面、目录、正文、图表和附录的版式。制作报告时，替换题目、署名、正文及图表内容。"},
        {"type": "heading", "level": 1, "text": "报告背景"},
        {"type": "paragraph", "text": "说明研究问题、任务背景与研究范围，明确报告要解决的实际问题。"},
        {"type": "heading", "level": 2, "text": "资料与方法"},
        {"type": "paragraph", "text": "记录资料来源、调查范围及处理方法，说明材料筛选与核查过程。"},
        {"type": "heading", "level": 3, "text": "资料整理"},
        {"type": "paragraph", "text": "按主题整理调研资料，为后续分析保留可追溯的依据。"},
        {"type": "heading", "level": 1, "text": "图表与结果"},
        {"type": "heading", "level": 2, "text": "表格示例"},
        {"type": "table", "headers": ["编号", "项目", "说明"], "rows": [
            ["DS-01", "访谈资料", "填写访谈对象、时间与记录编号。"],
            ["DS-02", "公开资料", "填写发布机构、资料标题与获取方式。"]
        ], "widths_cm": [2.3, 4.0, 9.3], "caption": "三线表排版示例"},
        {"type": "heading", "level": 2, "text": "图形示例"},
        {"type": "paragraph", "text": "图[[REF:layout_demo]]展示资料整理与呈现的排版示例。"},
        {"type": "figure", "path": str(ASSET_DIR / "layout-demo.png"), "caption": "图形与题注排版示例", "bookmark": "layout_demo", "width_cm": 14.0},
        {"type": "heading", "level": 1, "text": "参考文献", "numbered": False},
        {"type": "reference", "text": "参考文献条目：作者或机构、文献题名、文献类型、发表日期，以及适用的网址和访问日期。"},
        {"type": "heading", "level": 1, "text": "补充材料", "appendix": True},
        {"type": "paragraph", "text": "收录调查提纲、数据说明等与报告直接相关的补充材料。"},
    ]


def _string(value: Any, description: str, *, empty: bool = False) -> str:
    if not isinstance(value, str) or (not empty and not value.strip()):
        raise ValueError(f"{description} must be a {'possibly empty ' if empty else 'nonempty '}string")
    return value


def _width(value: Any, description: str) -> float:
    if isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(value) or value <= 0:
        raise ValueError(f"{description} must be a finite positive number")
    return float(value)


def validate_blocks(blocks: Any, base_dir: Path) -> tuple[list[dict[str, Any]], dict[str, int]]:
    """Validate the complete input before creating a destination file."""
    if not isinstance(blocks, list) or not blocks:
        raise ValueError("body JSON must contain a nonempty blocks array")
    normalized: list[dict[str, Any]] = []
    bookmarks: dict[str, int] = {}
    figure_number = 0
    for index, source in enumerate(blocks, start=1):
        if not isinstance(source, dict):
            raise ValueError(f"Block {index} must be an object")
        block = dict(source)
        kind = block.get("type", block.get("kind"))
        block["type"] = kind
        if kind == "heading":
            level = block.get("level")
            if isinstance(level, bool) or level not in (1, 2, 3):
                raise ValueError(f"Block {index}: heading level must be 1, 2 or 3")
            _string(block.get("text"), f"Block {index}: text")
            for option in ("numbered", "appendix"):
                if option in block and not isinstance(block[option], bool):
                    raise ValueError(f"Block {index}: {option} must be a boolean")
            if block.get("appendix") and not block.get("numbered", True):
                raise ValueError(f"Block {index}: an appendix heading must be numbered")
        elif kind in ("paragraph", "reference"):
            _string(block.get("text"), f"Block {index}: text", empty=(kind == "paragraph"))
        elif kind == "table":
            headers, rows = block.get("headers"), block.get("rows")
            if not isinstance(headers, list) or not headers:
                raise ValueError(f"Block {index}: table headers must be a nonempty array")
            for value in headers:
                _string(value, f"Block {index}: table header", empty=True)
            if not isinstance(rows, list) or any(not isinstance(row, list) or len(row) != len(headers) for row in rows):
                raise ValueError(f"Block {index}: each table row must have {len(headers)} cells")
            for row in rows:
                for value in row:
                    if not isinstance(value, (str, int, float)) or isinstance(value, bool):
                        raise ValueError(f"Block {index}: table cell must be a string or number")
                    if isinstance(value, float) and not math.isfinite(value):
                        raise ValueError(f"Block {index}: table numbers must be finite")
            block["rows"] = [[str(value) for value in row] for row in rows]
            if "widths_cm" in block:
                widths = block["widths_cm"]
                if not isinstance(widths, list) or len(widths) != len(headers):
                    raise ValueError(f"Block {index}: widths_cm must match the number of columns")
                block["widths_cm"] = [_width(value, f"Block {index}: column width") for value in widths]
                if sum(block["widths_cm"]) > TEXT_WIDTH_CM + 0.001:
                    raise ValueError(f"Block {index}: table widths exceed {TEXT_WIDTH_CM}cm")
            if "caption" in block:
                _string(block["caption"], f"Block {index}: caption", empty=True)
        elif kind == "figure":
            image_path = Path(_string(block.get("path"), f"Block {index}: path"))
            if not image_path.is_absolute():
                image_path = base_dir / image_path
            image_path = image_path.resolve()
            if not image_path.is_file():
                raise ValueError(f"Block {index}: image does not exist: {image_path}")
            with Image.open(image_path) as image:
                image.verify()
            block["path"] = image_path
            _string(block.get("caption"), f"Block {index}: caption")
            name = _string(block.get("bookmark"), f"Block {index}: bookmark")
            if not BOOKMARK_PATTERN.fullmatch(name):
                raise ValueError(f"Block {index}: bookmark must start with a letter and contain at most 40 letters/digits/underscores")
            if name in bookmarks:
                raise ValueError(f"Block {index}: duplicate figure bookmark: {name}")
            figure_number += 1
            bookmarks[name] = figure_number
            if "width_cm" in block:
                block["width_cm"] = _width(block["width_cm"], f"Block {index}: figure width")
                if block["width_cm"] > TEXT_WIDTH_CM:
                    raise ValueError(f"Block {index}: figure width exceeds {TEXT_WIDTH_CM}cm")
        else:
            raise ValueError(f"Block {index}: unsupported type {kind!r}")
        normalized.append(block)
    for index, block in enumerate(normalized, start=1):
        if block["type"] in ("paragraph", "reference"):
            for match in REF_PATTERN.finditer(block["text"]):
                if match.group(1) not in bookmarks:
                    raise ValueError(f"Block {index}: undefined figure bookmark {match.group(1)!r}")
            remaining = REF_PATTERN.sub("", block["text"])
            if "[[REF:" in remaining:
                raise ValueError(f"Block {index}: malformed [[REF:bookmark]] marker")
    return normalized, bookmarks


def build_document(metadata: ReportMetadata, blocks: list[dict[str, Any]],
                   bookmarks: dict[str, int], *, include_figure_list: bool = True) -> DocumentObject:
    """Build in memory; this helper never writes a DOCX to the filesystem."""
    document = Document()
    configure_styles(document)
    settings = document.settings.element
    xml_child(settings, "w:updateFields", val="true")
    xml_child(settings, "w:defaultTabStop", val="420")
    # Do not retain python-docx's stock author, timestamps or template comments.
    core = document.core_properties
    for child in list(core._element):
        core._element.remove(child)
    core.title = metadata.title
    if metadata.author:
        core.author = metadata.author
    if metadata.report_type:
        core.subject = metadata.report_type
    create_cover(document, metadata)
    create_front_matter(document, metadata, include_figure_list=bool(bookmarks) and include_figure_list)
    body = document.add_section(WD_SECTION_START.NEW_PAGE)
    configure_section(body, page_format="decimal")
    # Share only the front matter's established running header/footer.
    body.header.is_linked_to_previous = True
    body.footer.is_linked_to_previous = True
    title = document.add_paragraph(metadata.title, style="Report Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    numbering = add_numbering(document)
    appendix_numbering = add_numbering(document, appendix=True)
    appendix_active = False
    figure_number = 0
    table_number = 0
    for block in blocks:
        kind = block["type"]
        if kind == "heading":
            level = block["level"]
            paragraph = document.add_paragraph(block["text"], style=f"Heading {level}")
            if block.get("appendix"):
                appendix_active = True
            if block.get("numbered", True):
                apply_numbering(paragraph, appendix_numbering if appendix_active else numbering, level)
            else:
                # numId=0 explicitly suppresses numbering inherited from an outline style.
                apply_numbering(paragraph, 0, level)
        elif kind in ("paragraph", "reference"):
            paragraph = document.add_paragraph(style="Report Reference" if kind == "reference" else "Normal")
            format_paragraph(paragraph, first_line=0 if kind == "reference" else 24,
                             keep_together=(kind == "reference"))
            add_text_with_refs(paragraph, block["text"], bookmarks)
        elif kind == "table":
            if block.get("caption"):
                table_number += 1
            add_three_line_table(document, block["headers"], block["rows"], block.get("widths_cm"),
                                 caption=block.get("caption", ""), number=table_number)
        elif kind == "figure":
            figure_number += 1
            add_figure(document, block["path"], block["caption"], number=figure_number,
                       bookmark=block["bookmark"], bookmark_id=figure_number,
                       width_cm=block.get("width_cm"))
    return document


def save_new_document(document: DocumentObject, destination: Path) -> None:
    """Use exclusive creation so even a race cannot replace an existing file."""
    if destination.exists():
        raise FileExistsError(f"Refusing to overwrite existing file: {destination}")
    destination.parent.mkdir(parents=True, exist_ok=True)
    created = False
    try:
        with destination.open("xb") as stream:
            created = True
            document.save(stream)
    except Exception:
        if created:
            destination.unlink(missing_ok=True)
        raise


def make_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="Create a new editable CSU-style Word report.")
    parser.add_argument("--out", required=True, type=Path, help="New .docx destination; existing files are refused.")
    parser.add_argument("--title", required=True)
    parser.add_argument("--type", default="", dest="report_type")
    parser.add_argument("--subtitle", default="")
    parser.add_argument("--college", default="")
    parser.add_argument("--class-name", default="")
    parser.add_argument("--student-id", default="")
    parser.add_argument("--author", default="")
    parser.add_argument("--advisor", default="")
    parser.add_argument("--header-title", default="")
    parser.add_argument("--no-figure-list", action="store_true")
    parser.add_argument("--body-json", type=Path,
                        help='UTF-8 JSON object {"blocks": [...]}; image paths resolve relative to this file.')
    return parser


def main(argv: list[str] | None = None) -> int:
    parser = make_parser()
    args = parser.parse_args(argv)
    try:
        _string(args.title, "--title")
        out = args.out.expanduser().resolve()
        if out.suffix.lower() != ".docx":
            raise ValueError("--out must end in .docx")
        if out.exists():
            raise FileExistsError(f"Refusing to overwrite existing file: {out}")
        for name in ("csu_logo_full.png", "csu_seal.png"):
            if not (ASSET_DIR / name).is_file():
                raise ValueError(f"Required skill asset is missing: {ASSET_DIR / name}")
        if args.body_json:
            input_path = args.body_json.expanduser().resolve()
            payload = json.loads(input_path.read_text(encoding="utf-8-sig"))
            if not isinstance(payload, dict) or "blocks" not in payload:
                raise ValueError('body JSON must be an object containing "blocks"')
            raw_blocks = payload["blocks"]
            base_dir = input_path.parent
        else:
            raw_blocks = default_blocks()
            base_dir = ASSET_DIR
        blocks, bookmarks = validate_blocks(raw_blocks, base_dir)
        metadata = ReportMetadata(title=args.title, report_type=args.report_type, subtitle=args.subtitle,
                                  college=args.college, class_name=args.class_name,
                                  student_id=args.student_id, author=args.author,
                                  advisor=args.advisor, header_title=args.header_title)
        document = build_document(metadata, blocks, bookmarks, include_figure_list=not args.no_figure_list)
        save_new_document(document, out)
        print(str(out))
    except (ValueError, FileNotFoundError, FileExistsError, OSError, json.JSONDecodeError) as error:
        parser.error(str(error))
    return 0


if __name__ == "__main__":
    sys.exit(main())
