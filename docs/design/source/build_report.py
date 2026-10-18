"""Build Word and the matching XeLaTeX source from body.json; preserve editable fields."""
from pathlib import Path
import json,sys,importlib.util,re,shutil,argparse
ROOT=Path(__file__).resolve().parent
PARSER=argparse.ArgumentParser();PARSER.add_argument('--docx',action='store_true');PARSER.add_argument('--tex',action='store_true');args=PARSER.parse_args()
data=json.loads((ROOT/'body.json').read_text(encoding='utf8'));blocks=data['blocks'];meta=data['metadata']
for b in blocks:
 if b['type']=='figure':
  for k in ['path','pdf']:
   b[k]=str(ROOT/'figures'/Path(b[k]).name)
def esc(text):
 mapping={'\\':r'\textbackslash{}','&':r'\&','%':r'\%','$':r'\$','#':r'\#','_':r'\_','{':r'\{','}':r'\}','~':r'\textasciitilde{}','^':r'\textasciicircum{}'}
 return ''.join(mapping.get(c,c) for c in str(text))
def prose(text):
 # Natural break opportunities in code identifiers, paths, and SQL expressions.
 return ''.join(esc(c)+(r'\allowbreak{}' if c in '_/.,' else '') for c in str(text))
def textext(text):
 out=[];pos=0
 for m in re.finditer(r'\[\[REF:([\w]+)\]\]',text):out+=[prose(text[pos:m.start()]),r'\ref{'+m[1]+'}'];pos=m.end()
 out.append(prose(text[pos:]));return ''.join(out)
if args.docx or not args.tex:
 from docx.shared import Pt
 from docx.enum.text import WD_LINE_SPACING,WD_ALIGN_PARAGRAPH
 from docx.oxml.ns import qn
 word_skill=ROOT/'template/csu-word-report/scripts/new_report.py'
 if not word_skill.exists():word_skill=Path('D:/codex/skills/csu-word-report/scripts/new_report.py')
 spec=importlib.util.spec_from_file_location('csu_word',word_skill);mod=importlib.util.module_from_spec(spec);sys.modules['csu_word']=mod;spec.loader.exec_module(mod)
 validated,bookmarks=mod.validate_blocks(blocks,ROOT)
 doc=mod.build_document(mod.ReportMetadata(**meta),validated,bookmarks)
 for table in doc.tables[1:]:
  for ri,row in enumerate(table.rows):
   for cell in row.cells:
    # Only genuinely short IDs stay unbroken; field names and paths may wrap.
    if not re.fullmatch(r'(?:UC|BR|QR)-\d{2}',cell.text.strip()):
     for nw in list(cell._tc.get_or_add_tcPr().findall(qn('w:noWrap'))):nw.getparent().remove(nw)
    for para in cell.paragraphs:
     mod.format_paragraph(para,first_line=0,line=18,keep_next=ri==0,keep_together=True)
     for run in para.runs:mod.set_font(run,name=mod.HEADING_FONT if ri==0 else mod.BODY_FONT,size=10.5,bold=ri==0)
 for name in ['TOC 1','TOC 2','TOC 3']:
  if name in doc.styles:
   st=doc.styles[name];mod.set_font(st,size=11);st.paragraph_format.line_spacing=Pt(22);st.paragraph_format.line_spacing_rule=WD_LINE_SPACING.EXACTLY
 for para in doc.paragraphs:
  if para.text=='摘要':
   mod.format_paragraph(para,first_line=0,line=26,before=6,after=8,alignment=WD_ALIGN_PARAGRAPH.CENTER,keep_next=True)
   for run in para.runs:mod.set_font(run,name=mod.HEADING_FONT,size=14,bold=True)
 target=ROOT.parent/'安德鲁牙科连锁系统设计文档.docx';doc.save(target);print(json.dumps({'docx':str(target),'blocks':len(blocks)},ensure_ascii=False))
if args.tex or not args.docx:
 project=ROOT/'latex';project.mkdir(exist_ok=True);(project/'media').mkdir(exist_ok=True);(project/'figures').mkdir(exist_ok=True)
 academic_assets=ROOT/'template/csu-academic-report/assets'
 if not academic_assets.exists():academic_assets=Path('D:/codex/skills/csu-academic-report/assets')
 for name in ['csu_logo_full.png','csu_seal.png']:shutil.copyfile(academic_assets/name,project/'media'/name)
 template=(academic_assets/'report-template.tex').read_text(encoding='utf8')
 mapping={'ReportType':'','ReportTitle':meta['title'],'ReportSubtitle':'','ReportCollege':meta['college'],'ReportClass':meta['class_name'],'ReportStudentId':meta['student_id'],'ReportAuthor':meta['author'],'ReportAdvisor':meta['advisor'],'ReportHeaderTitle':meta['header_title']}
 for key,value in mapping.items():template=re.sub(r'\\newcommand\{\\'+key+r'\}\{[^\n]*\}',lambda _:"\\newcommand{\\"+key+'}{'+esc(value)+'}',template)
 pre=template[:template.index('% BODY START')]
 pre=pre.replace(r'\usepackage{graphicx}',r'\usepackage{graphicx}'+'\n'+r'\usepackage{needspace,float,seqsplit}')
 # Use TeX's page builder rather than a stale pagegoal after longtable output.
 pre=re.sub(r'\\newcommand\{\\ReportHeadingSpace\}\[1\]\{\\par.*?\}',lambda _:r'\newcommand{\ReportHeadingSpace}[1]{\par\Needspace{#1\baselineskip}}',pre,flags=re.S)
 pre=pre.replace(r'\sloppy',r'\sloppy'+'\n'+r'\setlength{\emergencystretch}{2em}\clubpenalty=10000\widowpenalty=10000\displaywidowpenalty=10000'+'\n'+r'\newcommand{\Cell}[1]{\begin{minipage}[t]{\linewidth}\raggedright\strut #1\strut\end{minipage}}')
 body=[r'\begin{center}{\sffamily\bfseries\fontsize{18pt}{26pt}\selectfont '+esc(meta['title'])+r'\par}\end{center}']
 appendix=False
 for bi,b in enumerate(blocks):
  typ=b['type']
  if typ=='heading':
   if b.get('appendix') and not appendix:body.append(r'\appendix');appendix=True
   cmd=['section','subsection','subsubsection'][b['level']-1]
   if bi+1<len(blocks) and blocks[bi+1]['type']=='heading':body.append(r'\Needspace{9\baselineskip}')
   elif bi+1<len(blocks) and blocks[bi+1]['type']=='table':body.append(r'\Needspace{7\baselineskip}')
   if b.get('numbered',True):body.append('\\'+cmd+'{'+esc(b['text'])+'}')
   else:body.append('\\'+cmd+'*{'+esc(b['text'])+'}\n'+r'\phantomsection\addcontentsline{toc}{'+cmd+'}{'+esc(b['text'])+'}')
  elif typ in ['paragraph','reference']:
   value=b['text']
   if value=='摘要':body.append(r'\begin{center}{\sffamily\bfseries\large 摘\quad 要}\end{center}');continue
   if typ=='reference':
    pieces=[];at=0
    for m in re.finditer(r'https?://[^\s，；]+',value):pieces+=[esc(value[at:m.start()]),r'\url{'+m[0]+'}'];at=m.end()
    pieces.append(esc(value[at:]));body.append(r'\noindent '+''.join(pieces)+r'\par');continue
   for line in value.split('\n'):body.append(textext(line)+r'\par')
  elif typ=='figure':
   src=Path(b['pdf']);shutil.copyfile(src,project/'figures'/src.name)
   body.append(r'\begin{figure}[H]\centering\includegraphics[width=\textwidth]{figures/'+src.name+'}\n'+r'\caption{'+esc(b['caption'])+'}'+r'\label{'+b['bookmark']+'}'+r'\end{figure}')
  elif typ=='table':
   # Long tables share one header and retain real text, with no vertical rules.
   widths=b['widths_cm'];cols='@{}'+'@{\\hspace{0.16cm}}'.join('p{'+f'{max(.5,w-.16):.3f}'+'cm}' for w in widths)+'@{}'
   def cell(value):return r'\Cell{'+r'\\'.join(prose(x) for x in str(value).split('\n'))+'}'
   header=' & '.join(r'\bfseries '+cell(x) for x in b['headers'])+r'\\\midrule'
   final_table=b.get('caption')=='交付及复现入口'
   leading='13' if final_table else '15'
   content=[r'{\fontsize{10.5pt}{'+leading+r'pt}\selectfont\setlength{\tabcolsep}{0pt}\renewcommand{\arraystretch}{1.08}',r'\begin{longtable}{'+cols+'}']
   if b.get('caption'):content.append(r'\caption{'+esc(b['caption'])+r'}\\')
   content.extend([r'\toprule',header,r'\endfirsthead',r'\toprule',header,r'\endhead',r'\bottomrule\endlastfoot'])
   for row in b['rows']:content.append(' & '.join(cell(x) for x in row)+(r'\\[1pt]' if final_table else r'\\[3pt]'))
   content.extend([r'\end{longtable}',r'}']);body+=content
 (project/'main.tex').write_text(pre+'\n'+'\n'.join(body)+'\n\\end{document}\n',encoding='utf8')
 print(json.dumps({'tex':str(project/'main.tex'),'figures':sum(b['type']=='figure' for b in blocks)},ensure_ascii=False))
